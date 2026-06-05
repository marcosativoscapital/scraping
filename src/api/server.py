"""Servidor FastAPI — backend do dashboard + integração com extensão Chrome."""

from __future__ import annotations

import contextvars
import csv
import io
import json
import logging
import os
import threading
from datetime import datetime
from pathlib import Path
from typing import Any, Optional
from uuid import uuid4

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, Header, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from ..claude_agent.classifier import classify_company
from ..claude_agent.client import GeminiClient
from ..claude_agent.personalize import generate_trigger
from ..claude_agent.scorer import score_lead
from ..db.store import Store
from ..workspaces.control import ControlStore
from ..workspaces.anamnese import analyze_anamnese
from ..enrichers.brasil_api import enrich_with_cnpj
from ..integrations.intercom import push_lead_to_intercom
from ..jobs.monitor import run_monitor
from ..jobs.rescore import run_rescore
from ..jobs.scheduler import Scheduler
from ..outbound.orchestrator import generate_and_store, send_outbound
from .auth import google_client_id, google_enabled, verify_google_id_token
from ..output.csv_writer import hydrate_db_row, write_leads_csv
from ..pipeline import run_pipeline
from ..playbooks.library import get_library
from ..playbooks.selector import select_playbooks_for_lead
from ..scrapers.linkedin import parse_linkedin_payload
from ..sdr.queue import SDRQueue
from ..claude_agent.client import GeminiClient
from ..enrichers.web_enricher import (
    discover_leads_via_icp,
    enrich_and_save,
    enrich_top_n,
    find_and_save_decisores,
)
from ..integrations.exact_spotter import push_and_mark as exact_push

load_dotenv()
logger = logging.getLogger(__name__)

API_TOKEN = os.environ.get("API_TOKEN", "change-me-in-production")
API_HOST = os.environ.get("API_HOST", "127.0.0.1")
API_PORT = int(os.environ.get("API_PORT", "8765"))

STATIC_DIR = Path(__file__).resolve().parent.parent / "dashboard" / "static"

app = FastAPI(
    title="Solve Scraper API",
    description="Backend do dashboard e extensão Chrome de prospecção da Solvefy.",
    version="0.2.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["chrome-extension://*", "http://localhost:*", "http://127.0.0.1:*"],
    allow_origin_regex=r"^(chrome-extension://.*|https?://(localhost|127\.0\.0\.1)(:\d+)?)$",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def _workspace_context(request, call_next):
    """Define o workspace ativo a partir do header X-Workspace-Id (default #1)."""
    tok = _CURRENT_WS.set(request.headers.get("x-workspace-id") or "1")
    try:
        return await call_next(request)
    finally:
        _CURRENT_WS.reset(tok)


# ====== STATE ======
JOBS: dict[str, dict[str, Any]] = {}

# Multi-tenant: cada workspace tem seu próprio banco. O STORE é um proxy que resolve
# o Store do workspace atual (header X-Workspace-Id) por requisição. Auth/sessões são
# globais e sempre usam o workspace #1 (AUTH_STORE).
CONTROL = ControlStore()
AUTH_STORE = CONTROL.get_data_store(1)
_CURRENT_WS: "contextvars.ContextVar[str]" = contextvars.ContextVar("ws_id", default="1")


class _WorkspaceStoreProxy:
    """Delega cada acesso para o Store do workspace atual."""

    def __getattr__(self, name):
        return getattr(CONTROL.get_data_store(_CURRENT_WS.get()), name)


STORE = _WorkspaceStoreProxy()


def ws_store(x_workspace_id: str | None = None) -> Store:
    """Store do workspace indicado (ou do contexto atual)."""
    return CONTROL.get_data_store(x_workspace_id if x_workspace_id is not None else _CURRENT_WS.get())


SCHED = Scheduler()
SDR = SDRQueue(STORE)


# ====== AUTH ======
def _auth(token: str | None) -> None:
    if token == API_TOKEN:
        return
    if AUTH_STORE.get_session(token):
        return
    raise HTTPException(status_code=401, detail="Token inválido")


def current_user(token: str | None) -> Optional[dict]:
    """Usuário da sessão Google (ou None se for o token único/extensão)."""
    if token and token != API_TOKEN:
        return AUTH_STORE.get_session(token)
    return None


class GoogleAuthPayload(BaseModel):
    credential: str


@app.get("/auth/config")
def auth_config():
    """Público: o frontend decide se mostra o 'Entrar com Google'."""
    return {"google_enabled": google_enabled(), "client_id": google_client_id()}


@app.post("/auth/google")
def auth_google(payload: GoogleAuthPayload):
    res = verify_google_id_token(payload.credential)
    if not res.get("ok"):
        raise HTTPException(401, res.get("erro") or "falha no login Google")
    token = AUTH_STORE.create_session(res["email"], res.get("nome"))
    AUTH_STORE.log_event("login", {"email": res["email"], "via": "google"})
    return {"ok": True, "token": token, "email": res["email"], "nome": res.get("nome")}


@app.get("/auth/me")
def auth_me(x_api_token: Optional[str] = Header(default=None)):
    u = current_user(x_api_token)
    if u:
        return {"autenticado": True, "email": u["email"], "nome": u.get("nome"), "via": "google"}
    if x_api_token == API_TOKEN:
        return {"autenticado": True, "email": None, "nome": "Token", "via": "token"}
    raise HTTPException(401, "Não autenticado")


@app.post("/auth/logout")
def auth_logout(x_api_token: Optional[str] = Header(default=None)):
    AUTH_STORE.delete_session(x_api_token)
    return {"ok": True}


# ====== WORKSPACES (multi-tenant) ======
class MemberPayload(BaseModel):
    email: str
    role: str = "leitor"


def _ws_public(w: dict) -> dict:
    out = {k: w.get(k) for k in ("id", "slug", "nome", "produto", "site", "descricao", "icp", "cor", "criado_em")}
    out["owner_email"] = w.get("owner_email")
    out["verticais"] = _lead_payload(w.get("anamnese_json")).get("verticais_sugeridas") or []
    return out


def _docx_text(raw: bytes) -> str:
    try:
        from docx import Document  # python-docx (opcional)
        return "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    except Exception:
        return ""


@app.get("/workspaces")
def workspaces_list(x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    email = (current_user(x_api_token) or {}).get("email")
    wss = CONTROL.list_workspaces(email) if email else CONTROL.list_workspaces()
    out = []
    for w in wss:
        item = _ws_public(w)
        item["role"] = CONTROL.role_of(w["id"], email) or "admin"  # dev/sem login = admin
        out.append(item)
    return {"workspaces": out}


@app.get("/workspaces/{ws_id:int}")
def workspaces_get(ws_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    w = CONTROL.get_workspace(ws_id)
    if not w:
        raise HTTPException(404, "Workspace não encontrado")
    anamnese = json.loads(w["anamnese_json"]) if w.get("anamnese_json") else None
    return {"workspace": _ws_public(w), "anamnese": anamnese, "membros": CONTROL.list_members(ws_id)}


@app.post("/workspaces")
async def workspaces_create(
    nome: str = Form(...),
    produto: str = Form(""),
    site: str = Form(""),
    descricao: str = Form(""),
    icp: str = Form(""),
    cor: str = Form("cpaas"),
    membros: str = Form("[]"),
    owner_email: str = Form(""),
    descricao_file: Optional[UploadFile] = File(None),
    icp_file: Optional[UploadFile] = File(None),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    # dono = usuário autenticado (login Google) ou, em modo token, o e-mail informado pelo front
    owner = (current_user(x_api_token) or {}).get("email") or (owner_email.strip() or None)

    anexos: list[tuple[bytes, str]] = []
    extra: list[str] = []
    for f in (descricao_file, icp_file):
        if not f:
            continue
        raw = await f.read()
        if not raw:
            continue
        mime = (f.content_type or "").lower()
        name = (f.filename or "").lower()
        if mime.startswith("text/") or name.endswith(".txt"):
            extra.append(f"[{f.filename}]\n" + raw.decode("utf-8", "ignore")[:8000])
        elif mime == "application/pdf" or name.endswith(".pdf") or mime.startswith("image/"):
            anexos.append((raw, mime or "application/pdf"))
        elif name.endswith(".docx"):
            t = _docx_text(raw)
            if t:
                extra.append(f"[{f.filename}]\n" + t[:8000])

    dados = {
        "nome": nome, "produto": produto, "site": site,
        "descricao": (descricao + (("\n\n" + "\n\n".join(extra)) if extra else "")).strip(),
        "icp": icp,
    }
    try:
        anamnese = analyze_anamnese(GeminiClient(), dados, anexos or None)
    except Exception as e:
        logger.exception("Anamnese falhou: %s", e)
        anamnese = {"erro": str(e)}

    try:
        membros_list = json.loads(membros) if membros else []
    except Exception:
        membros_list = []

    ws = CONTROL.create_workspace(
        nome, produto=produto, site=site, descricao=descricao, icp=icp, cor=cor,
        anamnese=anamnese, owner_email=owner, membros=membros_list,
    )
    try:
        CONTROL.get_data_store(ws["id"]).log_event("workspace_created", {"id": ws["id"], "nome": nome})
    except Exception:
        pass
    return {"ok": True, "workspace": _ws_public(ws), "anamnese": anamnese}


@app.get("/workspaces/{ws_id:int}/members")
def workspaces_members(ws_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    ws = CONTROL.get_workspace(ws_id)
    owner = (ws or {}).get("owner_email")
    membros = CONTROL.list_members(ws_id)
    # garante o dono na lista mesmo se não houver linha em workspace_members
    if owner and owner.lower() not in {(m.get("email") or "").lower() for m in membros}:
        membros = [{"email": owner, "role": "admin", "status": "active"}] + membros
    for m in membros:
        m["is_owner"] = bool(owner and (m.get("email") or "").lower() == owner.lower())
    membros.sort(key=lambda m: (not m.get("is_owner"),))  # dono primeiro
    return {"owner_email": owner, "membros": membros}


@app.post("/workspaces/{ws_id:int}/members")
def workspaces_add_member(ws_id: int, payload: MemberPayload, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    if not CONTROL.get_workspace(ws_id):
        raise HTTPException(404, "Workspace não encontrado")
    CONTROL.add_member(ws_id, payload.email, payload.role, status="invited")
    return {"ok": True, "membros": CONTROL.list_members(ws_id)}


@app.delete("/workspaces/{ws_id:int}")
def workspaces_delete(ws_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    if ws_id == 1:
        raise HTTPException(400, "O workspace principal não pode ser excluído")
    if not CONTROL.get_workspace(ws_id):
        raise HTTPException(404, "Workspace não encontrado")
    try:
        CONTROL.delete_workspace(ws_id)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True}


# ====== MODELS ======
class LinkedInPayload(BaseModel):
    source: str
    url: str
    items: list[dict[str, Any]]


class ScrapeRequest(BaseModel):
    vertical: str = Field(...)
    limit: Optional[int] = None
    enrich_email: bool = False


# ====== DASHBOARD ======
class NoCacheStaticFiles(StaticFiles):
    """Estático sempre revalidado: o navegador checa o ETag antes de reusar o cache,
    então nunca serve HTML/JS/CSS velho após um deploy (304 quando não mudou)."""

    async def get_response(self, path, scope):
        response = await super().get_response(path, scope)
        response.headers["Cache-Control"] = "no-cache, must-revalidate"
        return response


if STATIC_DIR.exists():
    app.mount("/dashboard", NoCacheStaticFiles(directory=str(STATIC_DIR), html=True), name="dashboard")


@app.get("/")
def root():
    return {
        "ok": True,
        "service": "Solve Scraper API",
        "version": "0.2.0",
        "jobs_ativos": len([j for j in JOBS.values() if j["status"] == "rodando"]),
        "dashboard": "/dashboard/",
    }


@app.get("/health")
def health():
    return {"status": "ok"}


# ====== STATS / LEADS / EVENTS (dashboard) ======
@app.get("/stats")
def stats(x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    return STORE.stats()


@app.get("/sales/cockpit")
def sales_cockpit(x_api_token: str | None = Header(default=None)):
    """Painel 'Meu dia' do vendedor: hoje, atrasadas, aguardando resposta, quentes, funil."""
    _auth(x_api_token)
    dia_de, dia_ate = _period_bounds("hoje")
    return STORE.cockpit(dia_de, dia_ate)


@app.get("/sales/reminders")
def sales_reminders(x_api_token: str | None = Header(default=None)):
    """Feed de lembretes acionáveis (sino): atrasadas, hoje, follow-ups, respostas, quentes."""
    _auth(x_api_token)
    dia_de, dia_ate = _period_bounds("hoje")
    return STORE.reminders(dia_de, dia_ate)


def _build_ics(a: dict) -> str:
    from datetime import datetime, timedelta

    def fmt(dt) -> str:
        return dt.strftime("%Y%m%dT%H%M%S")

    def esc(s) -> str:
        return str(s or "").replace("\\", "\\\\").replace("\n", "\\n").replace(",", "\\,").replace(";", "\\;")

    ini = (a.get("inicio_em") or "")[:16]
    try:
        start = datetime.fromisoformat(ini) if ini else datetime.now()
    except ValueError:
        start = datetime.now()
    end = start + timedelta(minutes=int(a.get("duracao_min") or 30))
    linhas = [
        "BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Solvefy//SolveScraper//PT", "BEGIN:VEVENT",
        f"UID:atividade-{a.get('id')}@solvescraper", f"DTSTART:{fmt(start)}", f"DTEND:{fmt(end)}",
        f"SUMMARY:{esc(a.get('titulo') or 'Atividade')}",
        f"LOCATION:{esc(a.get('cliente_empresa'))}",
        f"DESCRIPTION:{esc(a.get('descricao'))}",
        "END:VEVENT", "END:VCALENDAR",
    ]
    return "\r\n".join(linhas) + "\r\n"


@app.get("/atividades/{atividade_id:int}/calendar.ics")
def atividade_ics(atividade_id: int, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    a = STORE.get_atividade(atividade_id)
    if not a:
        raise HTTPException(404, "Atividade não encontrada")
    return StreamingResponse(
        iter([_build_ics(a)]),
        media_type="text/calendar",
        headers={"Content-Disposition": f"attachment; filename=atividade-{atividade_id}.ics"},
    )


@app.get("/db/leads")
def db_leads(
    vertical: str | None = Query(default=None),
    min_score: int = Query(default=0),
    limit: int = Query(default=200, le=2000),
    x_api_token: str | None = Header(default=None),
):
    _auth(x_api_token)
    leads = STORE.all_leads(vertical=vertical, min_score=min_score, limit=limit)
    return {"leads": leads, "count": len(leads)}


@app.get("/db/leads/{lead_id}")
def db_lead_detail(lead_id: int, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    with STORE.conn() as c:
        row = c.execute("SELECT * FROM leads WHERE id=?", (lead_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Lead não encontrado")
        outbound = c.execute(
            "SELECT * FROM outbound_messages WHERE lead_id=? ORDER BY canal", (lead_id,)
        ).fetchall()
    lead = dict(row)
    payload = _lead_payload(lead.get("payload_json"))
    return {
        "lead": lead,
        "outbound": [dict(o) for o in outbound],
        "atividades": STORE.list_atividades(lead_id=lead_id, order="asc"),
        "journey": payload.get("journey"),
        "enrichment": payload.get("web_enrichment"),
        "decisores": payload.get("decisores_extra") or [],
        "exact": payload.get("exact_spotter"),
    }


PIPELINE_STATUS_VALIDO = {"em_andamento", "ganho", "congelado", "perdido"}


class LeadPatch(BaseModel):
    pipeline_status: Optional[str] = None


@app.patch("/db/leads/{lead_id:int}")
def db_lead_update(lead_id: int, payload: LeadPatch, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    fields = payload.model_dump(exclude_unset=True)
    if not fields:
        raise HTTPException(400, "Nada para atualizar")
    if "pipeline_status" in fields and fields["pipeline_status"] not in PIPELINE_STATUS_VALIDO:
        raise HTTPException(400, f"pipeline_status inválido (use {sorted(PIPELINE_STATUS_VALIDO)})")
    if not STORE.update_lead_fields(lead_id, fields):
        raise HTTPException(404, "Lead não encontrado ou nada alterado")
    STORE.log_event("lead_pipeline_status", {"lead_id": lead_id, **fields})
    return {"ok": True, "lead_id": lead_id, **fields}


@app.get("/db/export.csv")
def db_export_csv(
    vertical: str | None = Query(default=None),
    min_score: int = Query(default=0),
    x_api_token: str | None = Header(default=None),
):
    _auth(x_api_token)
    leads = STORE.all_leads(vertical=vertical, min_score=min_score, limit=10000)
    if not leads:
        raise HTTPException(404, "Sem leads para exportar")

    buf = io.StringIO()
    cols = [
        "vertical", "segmento", "empresa", "cnpj", "razao_social", "site", "decisor_nome",
        "decisor_cargo", "decisor_linkedin", "email_provavel", "email_validado", "telefone",
        "porte_estimado", "score_icp", "recomendacao", "gatilho_personalizado", "observacoes",
        "fonte", "data_coleta", "criado_em", "atualizado_em",
    ]
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    w.writerows(hydrate_db_row(l) for l in leads)
    buf.seek(0)

    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=leads_{vertical or 'all'}.csv"},
    )


@app.get("/events")
def events(limit: int = Query(default=50), x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    return {"events": STORE.recent_events(limit=limit)}


# ====== LINKEDIN INGEST ======
@app.post("/linkedin/ingest")
def ingest_linkedin(payload: LinkedInPayload, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    leads_raw = parse_linkedin_payload(payload.model_dump())
    client = GeminiClient()
    processed = []

    for lead in leads_raw:
        try:
            cls = classify_company(client, lead)
            lead.update(
                {
                    "vertical": cls.get("vertical"),
                    "segmento": cls.get("segmento"),
                    "porte_estimado": cls.get("porte_estimado"),
                }
            )
            sc = score_lead(client, lead)
            lead["score_icp"] = sc.get("score")
            lead["recomendacao"] = sc.get("recomendacao")
            if (lead.get("score_icp") or 0) >= 50:
                lead["gatilho_personalizado"] = generate_trigger(client, lead)
            if lead.get("cnpj"):
                lead = enrich_with_cnpj(lead)
            lead["data_coleta"] = datetime.now().isoformat(timespec="seconds")
            STORE.upsert_lead(lead)
            STORE.log_event("novo_lead", {"empresa": lead.get("empresa"), "fonte": "linkedin"})
            processed.append(lead)
        except Exception as e:
            logger.exception("Falha em lead %s: %s", lead.get("empresa"), e)

    if processed:
        path = write_leads_csv(processed, vertical_tag="linkedin")
        return {
            "ok": True,
            "leads_processados": len(processed),
            "csv": str(path),
            "leads": processed,
        }
    return {"ok": True, "leads_processados": 0, "leads": []}


# ====== SCRAPE JOB ======
@app.post("/scrape")
def trigger_scrape(req: ScrapeRequest, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    job_id = str(uuid4())
    JOBS[job_id] = {
        "id": job_id,
        "status": "rodando",
        "vertical": req.vertical,
        "criado_em": datetime.now().isoformat(),
        "concluido_em": None,
        "leads": 0,
        "erro": None,
        "csv_path": None,
    }

    def _runner():
        try:
            leads = run_pipeline(
                vertical=req.vertical,
                limit=req.limit,
                enrich_email=req.enrich_email,
            )
            # Persiste no DB
            for lead in leads:
                STORE.upsert_lead(lead)
            STORE.log_event("scrape_completo", {"vertical": req.vertical, "leads": len(leads)})
            path = write_leads_csv(leads, vertical_tag=req.vertical) if leads else None
            JOBS[job_id].update(
                {
                    "status": "concluido",
                    "concluido_em": datetime.now().isoformat(),
                    "leads": len(leads),
                    "csv_path": str(path) if path else None,
                }
            )
        except Exception as e:
            logger.exception("Job %s falhou", job_id)
            JOBS[job_id].update(
                {
                    "status": "erro",
                    "erro": str(e),
                    "concluido_em": datetime.now().isoformat(),
                }
            )

    threading.Thread(target=_runner, daemon=True).start()
    return {"job_id": job_id, "status": "rodando"}


@app.get("/jobs/{job_id}")
def get_job(job_id: str, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    if job_id not in JOBS:
        raise HTTPException(404, "Job não encontrado")
    return JOBS[job_id]


@app.get("/jobs")
def list_jobs(x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    return list(JOBS.values())


@app.get("/outputs")
def list_outputs(x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    output_dir = Path("data/output")
    if not output_dir.exists():
        return []
    files = sorted(
        [
            {
                "arquivo": f.name,
                "caminho": str(f),
                "tamanho": f.stat().st_size,
                "modificado": datetime.fromtimestamp(f.stat().st_mtime).isoformat(),
            }
            for f in output_dir.glob("*.csv")
        ],
        key=lambda x: x["modificado"],
        reverse=True,
    )
    return files


# ====== MONITOR ======
@app.post("/monitor/{vertical}")
def monitor(vertical: str, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    try:
        return run_monitor(vertical, store=STORE)
    except Exception as e:
        logger.exception("Monitor falhou: %s", e)
        raise HTTPException(500, str(e))


# ====== RE-SCORE ======
@app.post("/rescore")
def rescore(
    days_since: int = Query(default=7),
    max_score: int = Query(default=70),
    x_api_token: str | None = Header(default=None),
):
    _auth(x_api_token)
    try:
        return run_rescore(days_since=days_since, max_score=max_score, store=STORE)
    except Exception as e:
        logger.exception("Rescore falhou: %s", e)
        raise HTTPException(500, str(e))


# ====== SCHEDULER ======
@app.post("/scheduler/start")
def sched_start(x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    SCHED.start()
    return SCHED.status()


@app.post("/scheduler/stop")
def sched_stop(x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    SCHED.stop()
    return SCHED.status()


@app.get("/scheduler/status")
def sched_status(x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    return SCHED.status()


# ====== OUTBOUND ======
@app.post("/outbound/generate/{lead_id}")
def outbound_generate(lead_id: int, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    with STORE.conn() as c:
        row = c.execute("SELECT * FROM leads WHERE id=?", (lead_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Lead não encontrado")
    lead = dict(row)
    try:
        messages = generate_and_store(lead_id, lead, store=STORE)
        return {"ok": True, "empresa": lead.get("empresa"), "messages": messages}
    except Exception as e:
        logger.exception("Geração de outbound falhou: %s", e)
        raise HTTPException(500, str(e))


class OutboundPatch(BaseModel):
    status: str  # aprovado | rejeitado


class ReplyPayload(BaseModel):
    msg_id: Optional[int] = None
    lead_id: Optional[int] = None
    nota: Optional[str] = None


@app.get("/outbound")
def outbound_list(
    status: Optional[str] = Query(default=None),
    limit: int = Query(default=200, le=1000),
    x_api_token: str | None = Header(default=None),
):
    _auth(x_api_token)
    msgs = STORE.all_outbound(status=status, limit=limit)
    return {"mensagens": msgs, "count": len(msgs)}


@app.patch("/outbound/{msg_id:int}")
def outbound_update(msg_id: int, payload: OutboundPatch, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    if payload.status not in ("aprovado", "rejeitado"):
        raise HTTPException(400, "status deve ser 'aprovado' ou 'rejeitado'")
    if not STORE.outbound_message(msg_id):
        raise HTTPException(404, "Mensagem não encontrada")
    STORE.update_outbound_status(msg_id, payload.status, erro=None)
    STORE.log_event("outbound_status", {"id": msg_id, "status": payload.status})
    return {"ok": True, "mensagem": STORE.outbound_message(msg_id)}


@app.post("/outbound/{msg_id:int}/send")
def outbound_send(msg_id: int, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    res = send_outbound(msg_id, store=STORE)
    if not res.get("ok"):
        raise HTTPException(400, res.get("erro") or "falha ao enviar")
    return {"ok": True, **res, "mensagem": STORE.outbound_message(msg_id)}


@app.post("/outbound/{msg_id:int}/reply")
def outbound_reply(msg_id: int, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    msg = STORE.outbound_message(msg_id)
    if not msg:
        raise HTTPException(404, "Mensagem não encontrada")
    now = datetime.now().isoformat(timespec="seconds")
    STORE.update_outbound_status(msg_id, "respondido", respondido_em=now)
    if msg.get("lead_id"):
        STORE.update_lead_fields(msg["lead_id"], {"sdr_status": "respondeu", "sdr_status_at": now})
    STORE.log_event("outbound_reply", {"id": msg_id, "lead_id": msg.get("lead_id")})
    return {"ok": True, "mensagem": STORE.outbound_message(msg_id)}


def _lead_payload(row_payload) -> dict:
    import json as _json
    if not row_payload:
        return {}
    return _json.loads(row_payload) if isinstance(row_payload, str) else (row_payload or {})


@app.post("/outbound/journey/{lead_id:int}")
def outbound_journey_generate(lead_id: int, x_api_token: str | None = Header(default=None)):
    """Gera (via Gemini) a jornada de contato ideal do lead e salva em payload_json."""
    _auth(x_api_token)
    import json as _json
    with STORE.conn() as c:
        row = c.execute("SELECT * FROM leads WHERE id=?", (lead_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Lead não encontrado")
    lead = dict(row)
    try:
        from ..claude_agent.journey import generate_journey
        journey = generate_journey(GeminiClient(), lead)
    except Exception as e:
        logger.exception("Geração de jornada falhou: %s", e)
        raise HTTPException(500, str(e))
    with STORE.conn() as c:
        pd = _lead_payload(lead.get("payload_json"))
        pd["journey"] = journey
        c.execute(
            "UPDATE leads SET payload_json=? WHERE id=?",
            (_json.dumps(pd, ensure_ascii=False, default=str), lead_id),
        )
    STORE.log_event("journey_generated", {"lead_id": lead_id, "empresa": lead.get("empresa")})
    return {"ok": True, "lead_id": lead_id, "empresa": lead.get("empresa"), "journey": journey}


@app.get("/outbound/journey/{lead_id:int}")
def outbound_journey_get(lead_id: int, x_api_token: str | None = Header(default=None)):
    """Retorna a jornada já gerada (ou null)."""
    _auth(x_api_token)
    with STORE.conn() as c:
        row = c.execute("SELECT payload_json FROM leads WHERE id=?", (lead_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Lead não encontrado")
    return {"lead_id": lead_id, "journey": _lead_payload(row["payload_json"]).get("journey")}


@app.post("/webhook/outbound/reply")
def outbound_webhook_reply(payload: ReplyPayload, x_api_token: str | None = Header(default=None)):
    """Webhook para CPaaS/parser futuro marcar resposta por msg_id ou lead_id."""
    _auth(x_api_token)
    now = datetime.now().isoformat(timespec="seconds")
    lead_id = payload.lead_id
    if payload.msg_id:
        STORE.update_outbound_status(payload.msg_id, "respondido", respondido_em=now)
        m = STORE.outbound_message(payload.msg_id)
        if m and m.get("lead_id"):
            lead_id = m["lead_id"]
    if lead_id:
        STORE.update_lead_fields(lead_id, {"sdr_status": "respondeu", "sdr_status_at": now})
    STORE.log_event("outbound_reply_webhook", {"msg_id": payload.msg_id, "lead_id": lead_id})
    return {"ok": True}


# ====== INTERCOM ======
@app.post("/intercom/push/{lead_id}")
def intercom_push(lead_id: int, x_api_token: str | None = Header(default=None)):
    _auth(x_api_token)
    with STORE.conn() as c:
        row = c.execute("SELECT * FROM leads WHERE id=?", (lead_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Lead não encontrado")
    result = push_lead_to_intercom(dict(row))
    STORE.log_event("intercom_push", {"lead_id": lead_id, "result": result})
    return result


@app.post("/intercom/push_batch")
def intercom_push_batch(
    min_score: int = Query(default=70),
    x_api_token: str | None = Header(default=None),
):
    _auth(x_api_token)
    leads = STORE.all_leads(min_score=min_score, limit=500)
    from ..integrations.intercom import push_batch
    result = push_batch(leads)
    STORE.log_event("intercom_batch", result)
    return result


# ====== PLAYBOOKS ======
@app.get("/playbooks")
def list_playbooks(x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    lib = get_library()
    return {
        "playbooks": [p.as_dict() for p in lib.all()],
        "objecoes": [{"id": o.id, "titulo": o.titulo, "resposta": o.resposta} for o in lib.objecoes],
    }


@app.get("/leads/{lead_id}/playbooks")
def lead_playbooks(lead_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    pbs = SDR.playbooks_for_lead(lead_id)
    lib = get_library()
    enriched = []
    for pb in pbs:
        full = lib.get(pb.get("playbook_id"))
        if full:
            d = full.as_dict()
            d.update(pb)
            enriched.append(d)
        else:
            enriched.append(pb)
    return {"lead_id": lead_id, "playbooks": enriched}


@app.post("/leads/{lead_id}/playbooks/regenerate")
def regenerate_playbooks(lead_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    with STORE.conn() as c:
        row = c.execute("SELECT * FROM leads WHERE id=?", (lead_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Lead não encontrado")
    lead = dict(row)
    try:
        pbs = select_playbooks_for_lead(GeminiClient(), lead, n=3)
        SDR.assign_playbooks(lead_id, pbs)
        return {"ok": True, "playbooks": pbs}
    except Exception as e:
        logger.exception("Falha ao regenerar playbooks: %s", e)
        raise HTTPException(500, str(e))


@app.post("/leads/{lead_id}/playbooks/{playbook_id}/status")
def set_playbook_status(
    lead_id: int,
    playbook_id: str,
    status: str = Query(..., description="sugerido|em_execucao|concluido|abandonado"),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    SDR.update_playbook_status(lead_id, playbook_id, status)
    return {"ok": True}


# ====== SDR ======
@app.get("/sdr/queue")
def sdr_queue(
    sdr_email: Optional[str] = Query(default=None),
    status: Optional[str] = Query(default=None),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    return {"queue": SDR.queue_for(sdr_email=sdr_email, status=status)}


@app.post("/sdr/assign")
def sdr_assign(
    lead_id: int = Query(...),
    sdr_email: str = Query(...),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    SDR.assign_lead(lead_id, sdr_email)
    return {"ok": True}


@app.post("/sdr/auto-assign")
def sdr_auto_assign(
    sdr_email: str = Query(...),
    min_score: int = Query(default=60),
    max_n: int = Query(default=20),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    n = SDR.auto_assign_hot_leads(sdr_email, min_score=min_score, max_n=max_n)
    return {"ok": True, "leads_atribuidos": n}


class ActivityPayload(BaseModel):
    lead_id: int
    sdr_email: str
    tipo: str
    canal: Optional[str] = None
    playbook_id: Optional[str] = None
    outcome: Optional[str] = None
    notas: Optional[str] = None


@app.post("/sdr/activity")
def sdr_log_activity(payload: ActivityPayload, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    aid = SDR.log_activity(**payload.model_dump())
    return {"ok": True, "activity_id": aid}


@app.get("/sdr/activities/{lead_id}")
def sdr_activities(lead_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    return {"activities": SDR.activities_for_lead(lead_id)}


@app.get("/sdr/metrics")
def sdr_metrics(
    sdr_email: Optional[str] = Query(default=None),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    return SDR.metrics(sdr_email=sdr_email)


ENRICH_JOBS: dict[str, dict[str, Any]] = {}


@app.post("/enrichment/top")
def enrich_top(
    n: int = Query(default=50),
    min_score: int = Query(default=60),
    x_api_token: Optional[str] = Header(default=None),
):
    """Dispara enrichment web em background dos top N leads."""
    _auth(x_api_token)
    job_id = str(uuid4())
    ENRICH_JOBS[job_id] = {
        "id": job_id,
        "status": "rodando",
        "iniciado_em": datetime.now().isoformat(),
        "concluido_em": None,
        "n": n,
        "enriquecidos": 0,
        "erros": 0,
        "resultados": [],
    }

    def _runner():
        try:
            result = enrich_top_n(n=n, min_score=min_score, store=STORE)
            ENRICH_JOBS[job_id].update({
                "status": "concluido",
                "concluido_em": datetime.now().isoformat(),
                "enriquecidos": result["enriquecidos"],
                "erros": result["erros"],
                "resultados": result["detalhes"],
            })
        except Exception as e:
            logger.exception("Job enrichment %s falhou", job_id)
            ENRICH_JOBS[job_id].update({
                "status": "erro",
                "erro": str(e),
                "concluido_em": datetime.now().isoformat(),
            })

    threading.Thread(target=_runner, daemon=True).start()
    return {"job_id": job_id, "status": "rodando"}


@app.get("/enrichment/jobs/{job_id}")
def enrich_job_status(job_id: str, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    if job_id not in ENRICH_JOBS:
        raise HTTPException(404, "Job não encontrado")
    return ENRICH_JOBS[job_id]


@app.post("/enrichment/lead/{lead_id}")
def enrich_lead(lead_id: int, x_api_token: Optional[str] = Header(default=None)):
    """Enriquece um lead único sob demanda."""
    _auth(x_api_token)
    return enrich_and_save(lead_id, store=STORE)


@app.post("/enrichment/lead/{lead_id}/decisores")
def enrich_lead_decisores(lead_id: int, x_api_token: Optional[str] = Header(default=None)):
    """Busca na web/LinkedIn o máximo de possíveis decisores da empresa do lead."""
    _auth(x_api_token)
    try:
        return find_and_save_decisores(lead_id, store=STORE)
    except Exception as e:  # noqa: BLE001
        logger.exception("Busca de decisores falhou (lead %s)", lead_id)
        raise HTTPException(500, f"Falha ao buscar decisores: {e}")


class DiscoverRequest(BaseModel):
    vertical: Optional[str] = None
    limit: int = 8
    refino: Optional[str] = None


@app.post("/leads/discover")
def leads_discover(
    req: DiscoverRequest,
    x_api_token: Optional[str] = Header(default=None),
    x_workspace_id: Optional[str] = Header(default=None),
):
    """Descobre leads reais (Gemini + Google Search) que batem com o ICP do workspace atual."""
    _auth(x_api_token)
    try:
        ws_id = int(x_workspace_id) if x_workspace_id is not None else 1
    except (ValueError, TypeError):
        ws_id = 1
    ws = CONTROL.get_workspace(ws_id)
    if not ws:
        raise HTTPException(404, "Workspace não encontrado")
    limit = max(1, min(int(req.limit or 8), 15))
    try:
        return discover_leads_via_icp(CONTROL.get_data_store(ws_id), ws, vertical_label=req.vertical, limit=limit, refino=req.refino)
    except Exception as e:  # noqa: BLE001
        logger.exception("Descoberta de leads falhou (ws %s)", ws_id)
        raise HTTPException(500, f"Falha na descoberta: {e}")


@app.post("/leads/{lead_id:int}/exact")
def leads_to_exact(lead_id: int, x_api_token: Optional[str] = Header(default=None)):
    """Envia o lead ao Exact Spotter (dry-run por padrão; loga o payload sem enviar)."""
    _auth(x_api_token)
    try:
        return exact_push(lead_id, store=STORE)
    except Exception as e:  # noqa: BLE001
        logger.exception("Envio ao Exact Spotter falhou (lead %s)", lead_id)
        raise HTTPException(500, f"Falha no envio ao Exact Spotter: {e}")


# ====== ATIVIDADES (vendas / oportunidades) ======


class AtividadePayload(BaseModel):
    lead_id: Optional[int] = None
    titulo: Optional[str] = None
    natureza: str = "evento"            # evento | tarefa | lembrete
    tipo: Optional[str] = None          # ligacao | videochamada | email | visita | almoco | personalizado
    inicio_em: Optional[str] = None     # ISO8601
    duracao_min: Optional[int] = None
    dia_inteiro: bool = False
    repeticao: str = "nenhuma"
    temperatura: Optional[str] = None   # muito_quente | quente | frio | muito_frio
    pipeline: str = "potencial_cliente"
    responsavel: Optional[str] = None
    contato_nome: Optional[str] = None
    descricao: Optional[str] = None
    tags: Optional[list[str]] = None
    status: str = "a_fazer"


class AtividadePatch(BaseModel):
    lead_id: Optional[int] = None
    titulo: Optional[str] = None
    natureza: Optional[str] = None
    tipo: Optional[str] = None
    inicio_em: Optional[str] = None
    duracao_min: Optional[int] = None
    dia_inteiro: Optional[bool] = None
    repeticao: Optional[str] = None
    temperatura: Optional[str] = None
    pipeline: Optional[str] = None
    responsavel: Optional[str] = None
    contato_nome: Optional[str] = None
    descricao: Optional[str] = None
    tags: Optional[list[str]] = None
    status: Optional[str] = None


def _period_bounds(periodo: Optional[str]) -> tuple[Optional[str], Optional[str]]:
    """Converte 'hoje'/'semana'/'mes' em limites ISO [início, fim) para inicio_em."""
    from datetime import timedelta

    if not periodo or periodo == "todos":
        return None, None
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    if periodo == "hoje":
        start, end = today, today + timedelta(days=1)
    elif periodo == "semana":
        start = today - timedelta(days=today.weekday())
        end = start + timedelta(days=7)
    elif periodo == "mes":
        start = today.replace(day=1)
        end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    else:
        return None, None
    return start.isoformat(), end.isoformat()


@app.get("/atividades")
def atividades_list(
    responsavel: Optional[str] = Query(default=None),
    periodo: Optional[str] = Query(default=None),
    tipo: Optional[str] = Query(default=None),
    temperatura: Optional[str] = Query(default=None),
    pipeline: Optional[str] = Query(default=None),
    status: Optional[str] = Query(default=None),
    lead_id: Optional[int] = Query(default=None),
    limit: int = Query(default=500, le=2000),
    order: str = Query(default="asc"),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    inicio_de, inicio_ate = _period_bounds(periodo)
    items = STORE.list_atividades(
        responsavel=responsavel,
        tipo=tipo,
        temperatura=temperatura,
        pipeline=pipeline,
        status=status,
        lead_id=lead_id,
        inicio_de=inicio_de,
        inicio_ate=inicio_ate,
        limit=limit,
        order=order,
    )
    return {"atividades": items, "count": len(items)}


@app.post("/atividades")
def atividades_create(payload: AtividadePayload, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    data = payload.model_dump()
    if not data.get("responsavel"):
        u = current_user(x_api_token)
        if u:
            data["responsavel"] = u["email"]
    aid = STORE.create_atividade(data)
    STORE.log_event(
        "atividade_criada",
        {"atividade_id": aid, "titulo": payload.titulo, "pipeline": payload.pipeline},
    )
    return {"ok": True, "id": aid, "atividade": STORE.get_atividade(aid)}


@app.get("/atividades/{atividade_id:int}")
def atividades_detail(atividade_id: int, x_api_token: Optional[str] = Header(default=None)):
    _auth(x_api_token)
    atv = STORE.get_atividade(atividade_id)
    if not atv:
        raise HTTPException(404, "Atividade não encontrada")
    return {"atividade": atv}


@app.patch("/atividades/{atividade_id:int}")
def atividades_update(
    atividade_id: int,
    payload: AtividadePatch,
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    fields = payload.model_dump(exclude_unset=True)
    if not fields:
        raise HTTPException(400, "Nada para atualizar")
    if not STORE.update_atividade(atividade_id, fields):
        raise HTTPException(404, "Atividade não encontrada ou nada alterado")
    if "status" in fields or "pipeline" in fields:
        STORE.log_event(
            "atividade_status",
            {"atividade_id": atividade_id, **{k: fields[k] for k in ("status", "pipeline") if k in fields}},
        )
    return {"ok": True, "atividade": STORE.get_atividade(atividade_id)}


def _range_bounds(ref: Optional[str], escala: str) -> tuple[str, str]:
    """Limites ISO [início, fim) para a janela de tempo (mes/semana/trimestre/ano)."""
    from datetime import timedelta

    base = datetime.now()
    if ref:
        try:
            base = datetime.fromisoformat(ref[:10])
        except ValueError:
            pass
    base = base.replace(hour=0, minute=0, second=0, microsecond=0)
    if escala == "semana":
        start = base - timedelta(days=(base.weekday() + 1) % 7)  # semana começa no domingo
        end = start + timedelta(days=7)
    elif escala == "trimestre":
        qm = ((base.month - 1) // 3) * 3 + 1
        start = base.replace(month=qm, day=1)
        nm, ny = qm + 3, base.year
        if nm > 12:
            nm, ny = nm - 12, ny + 1
        end = start.replace(year=ny, month=nm, day=1)
    elif escala == "ano":
        start = base.replace(month=1, day=1)
        end = start.replace(year=start.year + 1)
    else:  # mes
        start = base.replace(day=1)
        end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    return start.isoformat(), end.isoformat()


@app.get("/atividades/calendario")
def atividades_calendario(
    ref: Optional[str] = Query(default=None),
    escala: str = Query(default="mes"),
    responsavel: Optional[str] = Query(default=None),
    tipo: Optional[str] = Query(default=None),
    temperatura: Optional[str] = Query(default=None),
    pipeline: Optional[str] = Query(default=None),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    inicio, fim = _range_bounds(ref, "semana" if escala == "semana" else "mes")
    items = STORE.list_atividades(
        responsavel=responsavel, tipo=tipo, temperatura=temperatura, pipeline=pipeline,
        inicio_de=inicio, inicio_ate=fim, limit=2000,
    )
    return {"escala": escala, "inicio": inicio, "fim": fim, "atividades": items}


@app.get("/atividades/timeline")
def atividades_timeline(
    ref: Optional[str] = Query(default=None),
    escala: str = Query(default="mes"),
    responsavel: Optional[str] = Query(default=None),
    tipo: Optional[str] = Query(default=None),
    temperatura: Optional[str] = Query(default=None),
    pipeline: Optional[str] = Query(default=None),
    x_api_token: Optional[str] = Header(default=None),
):
    _auth(x_api_token)
    esc = escala if escala in ("mes", "trimestre", "ano") else "mes"
    inicio, fim = _range_bounds(ref, esc)
    items = STORE.list_atividades(
        responsavel=responsavel, tipo=tipo, temperatura=temperatura, pipeline=pipeline,
        inicio_de=inicio, inicio_ate=fim, limit=5000,
    )

    grupos: dict[int, dict] = {}
    sem_lead: list[dict] = []
    for a in items:
        rec = {
            "id": a["id"], "inicio_em": a["inicio_em"], "tipo": a["tipo"],
            "temperatura": a["temperatura"], "titulo": a["titulo"], "pipeline": a["pipeline"],
        }
        lid = a.get("lead_id")
        if not lid:
            sem_lead.append(rec)
            continue
        g = grupos.setdefault(lid, {
            "lead_id": lid,
            "empresa": a.get("cliente_empresa") or f"Lead #{lid}",
            "status": a.get("cliente_status"),
            "atividades": [],
        })
        g["atividades"].append(rec)
        if a.get("pipeline") == "pos_venda" and not g["status"]:
            g["status"] = "ganho"

    oportunidades = []
    for g in grupos.values():
        datas = sorted([x["inicio_em"][:10] for x in g["atividades"] if x["inicio_em"]])
        ciclo = 0
        if len(datas) >= 2:
            try:
                ciclo = (datetime.fromisoformat(datas[-1]) - datetime.fromisoformat(datas[0])).days
            except ValueError:
                ciclo = 0
        g["ciclo_dias"] = ciclo
        g["status"] = g["status"] or "em_andamento"
        oportunidades.append(g)
    oportunidades.sort(key=lambda x: x["empresa"].lower())

    return {"escala": esc, "inicio": inicio, "fim": fim, "oportunidades": oportunidades, "sem_lead": sem_lead}


def start():
    """Entry point: python -m src.api.server"""
    import uvicorn

    uvicorn.run(
        "src.api.server:app",
        host=API_HOST,
        port=API_PORT,
        reload=False,
    )


if __name__ == "__main__":
    start()
